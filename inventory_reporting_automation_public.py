from __future__ import annotations

from collections import Counter, defaultdict
from math import floor, isfinite
from fractions import Fraction
from io import BytesIO
from decimal import Decimal, InvalidOperation
from pathlib import Path
from datetime import datetime
import re
import sys

# ------------------------------------------------------------
# سازگاری خروجی فارسی با PowerShell / VS Code / Code Runner
# ------------------------------------------------------------
try:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    if hasattr(sys.stderr, "reconfigure"):
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    raise SystemExit(
        "کتابخانه openpyxl نصب نیست.\n"
        "در Command Prompt اجرا کن:\n"
        "pip install openpyxl"
    )


# ============================================================
# تنظیمات
# ============================================================
# پیشنهاد: فایل پایتون و دیکشنری کنار هم باشند و فایل‌های نماینده
# داخل پوشه inventories قرار بگیرند.
BASE_DIR = Path(__file__).resolve().parent

DICTIONARY_BASENAME = "Dictionary_v2_matched_updated"
INPUT_FOLDER = BASE_DIR / "inventories"
OUTPUT_FILE = BASE_DIR / "inventory_report.xlsx"

# فقط عنوان ستون‌هایی که در نمونه‌های تأییدشده دیده‌ایم.
# اگر بعداً نماینده‌ای عنوان متفاوتی برای کد کالا داشت،
# فقط همان عنوان را به این مجموعه اضافه کن.
CODE_COLUMN_CANDIDATES = {
    "کد کالا",
    "کدکالا",   # مثل فایل کرمان: «كدكالا» که بعد از نرمال‌سازی می‌شود «کدکالا»
    "کد",
}

# برنامه برای پیدا کردن ردیف هدر تا این تعداد ردیف اول هر شیت را بررسی می‌کند.
HEADER_SCAN_LIMIT = 20

# نام ستون‌های اصلی دیکشنری
D_REP = "نماینده"
D_REP_CODE = "کد کالای نماینده"
D_REP_PRODUCT = "نام کالای نماینده"
D_COMPANY_CODE = "کد شرکت"
D_COMPANY_NAME = "نام شرکت"
D_OWN_CODE = "کد کالای خودمان"
D_SHORT_NAME = "نام خلاصه"
D_BRAND = "برند"
D_PRODUCT_GROUP = "گروه کالا"
D_FLAVOR = "طعم"
D_WEIGHT = "وزن"
D_LEVEL = "سطح موجودی نماینده"
D_INPUT_COLUMN = "ستون مورد نیاز"
D_MIDDLE_PER_MASTER = "تعداد میدل در کارتن مادر"
D_UNITS_PER_CARTON = "تعداد در کارتن"


# ============================================================
# توابع کمکی
# ============================================================

PERSIAN_ARABIC_DIGITS = str.maketrans(
    "۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩",
    "01234567890123456789"
)


def normalize_text(value) -> str:
    """
    فقط برای مقایسه عنوان ستون‌ها استفاده می‌شود.
    فاصله‌های اضافه و تفاوت ی/ک عربی و فارسی را یکسان می‌کند.
    """
    if value is None:
        return ""

    text = str(value)
    text = text.replace("\ufeff", "")
    text = text.replace("\u200c", " ")
    text = text.replace("\u200f", "")
    text = text.replace("ي", "ی").replace("ى", "ی").replace("ك", "ک")
    text = re.sub(r"\s+", " ", text).strip()
    return text


# ترتیب صریح حروف فارسی برای مرتب‌سازی از الف تا ی.
PERSIAN_ALPHABET = "اآبپتثجچحخدذرزژسشصضطظعغفقکگلمنوهی"
PERSIAN_ORDER = {char: index for index, char in enumerate(PERSIAN_ALPHABET)}


def persian_sort_key(value):
    """
    کلید مرتب‌سازی فارسی مستقل از Locale ویندوز.
    فاصله و علائم نگارشی در ترتیب حروف نادیده گرفته می‌شوند.
    """
    text = normalize_text(value)
    text = (
        text.replace("أ", "ا")
        .replace("إ", "ا")
        .replace("ٱ", "ا")
        .replace("ؤ", "و")
        .replace("ئ", "ی")
        .replace("ۀ", "ه")
        .replace("ة", "ه")
    )

    key = []
    for char in text:
        if char in PERSIAN_ORDER:
            key.append((0, PERSIAN_ORDER[char]))
        elif char.isdigit():
            key.append((1, ord(char)))
        elif char.isspace() or char in "-_/()[]{}.,،؛:":
            continue
        else:
            # حروف لاتین یا سایر کاراکترها بعد از حروف فارسی قرار می‌گیرند.
            key.append((2, ord(char)))

    return tuple(key)


def to_decimal(value) -> Decimal:
    """
    تبدیل دقیق مقدار اکسل به Decimal برای اینکه جمع چند کد قبل از گرد کردن
    خطای اعشاری شناور ایجاد نکند.
    """
    if value is None:
        raise ValueError("مقدار خالی است")

    if isinstance(value, bool):
        raise ValueError("مقدار True/False عدد موجودی نیست")

    if isinstance(value, Decimal):
        number = value
    elif isinstance(value, int):
        number = Decimal(value)
    elif isinstance(value, float):
        if not isfinite(value):
            raise ValueError("عدد معتبر نیست")
        number = Decimal(str(value))
    else:
        text = str(value).strip().translate(PERSIAN_ARABIC_DIGITS)
        text = text.replace(",", "").replace("٬", "").replace(" ", "")
        if text == "":
            raise ValueError("مقدار خالی است")
        try:
            number = Decimal(text)
        except InvalidOperation:
            raise ValueError(f"مقدار عددی نیست: {value!r}")

    if not number.is_finite():
        raise ValueError("عدد معتبر نیست")

    return number


def normalize_weight(value):
    """
    وزن را برای نمایش در هدر آماده می‌کند.
    """
    if value is None or str(value).strip() == "":
        return None

    try:
        number = to_decimal(value)
    except ValueError:
        return None

    if number <= 0:
        return None

    if number == number.to_integral_value():
        return str(int(number))

    return format(number.normalize(), "f")


def representative_match_key(value) -> str:
    """
    برای تطبیق نام فایل نماینده:
    بیرجند فراقی == بیرجند-فراقی
    """
    text = normalize_text(value)
    text = re.sub(r"[\s\-_–—/]+", "", text)
    return text


def clean_flavor_weight_suffix(flavor: str, weight: str | None) -> str:
    """
    اگر وزن انتهای نام طعم آمده باشد، برای تصمیم‌گیری درباره‌ی نمایش وزن
    آن را موقتاً حذف می‌کند.
    """
    flavor = normalize_text(flavor)

    if not flavor or not weight:
        return flavor

    flavor_ascii = flavor.translate(PERSIAN_ARABIC_DIGITS)

    cleaned = re.sub(
        r"\s+\d+(?:\.\d+)?\s*(?:گرم|گرمی|g|gr)?\s*$",
        "",
        flavor_ascii,
        flags=re.IGNORECASE,
    ).strip()

    return cleaned or flavor


def section_and_base_flavor(entry: dict):
    """
    خروجی هرم:
      برند -> گروه کالا -> سکشن -> طعم

    Project-specific product/section rules are intentionally omitted
    from the public portfolio version.

    هر چیزی که «تک طعم» داشته باشد در همان سکشن با عنوان «تک طعم»
    تجمیع می‌شود.
    """
    brand = normalize_text(entry.get("brand"))
    group = normalize_text(entry.get("product_group"))
    flavor = normalize_text(entry.get("flavor"))
    weight = normalize_weight(entry.get("weight"))

    section = ""

    # Company/product-specific section rules have been removed for the
    # public portfolio version.

    if "تک طعم" in flavor:
        base_flavor = "تک طعم"
    else:
        base_flavor = clean_flavor_weight_suffix(flavor, weight)

    if not base_flavor:
        base_flavor = group

    return section, base_flavor


def output_key_from_entry(entry: dict):
    section, base_flavor = section_and_base_flavor(entry)

    display_flavor = (
        normalize_text(entry.get("display_flavor"))
        or base_flavor
    )

    return (
        normalize_text(entry.get("brand")),
        normalize_text(entry.get("product_group")),
        section,
        display_flavor,
    )


def flavor_sort_key(value):
    """
    ترتیب الفبایی فارسی، با یک استثنا:
    هر عنوانی که «جور» داشته باشد آخر همان سکشن قرار می‌گیرد.
    """
    text = normalize_text(value)
    is_assorted = 1 if "جور" in text else 0
    return (is_assorted, persian_sort_key(text))


def resolve_representative(file_stem: str, dictionary_data: dict):
    """
    ابتدا تطبیق دقیق، سپس تطبیق با حذف فاصله/خط تیره.
    اگر بیش از یک گزینه وجود داشته باشد، برنامه حدس نمی‌زند.
    """
    if file_stem in dictionary_data["representatives"]:
        return file_stem

    key = representative_match_key(file_stem)
    matches = dictionary_data["representative_aliases"].get(key, set())

    if len(matches) == 1:
        return next(iter(matches))

    if len(matches) > 1:
        raise ValueError(
            "نام فایل به بیش از یک نماینده می‌خورد: "
            + " | ".join(sorted(matches, key=persian_sort_key))
        )

    raise ValueError(
        "نام فایل با هیچ نماینده‌ای در دیکشنری تطبیق پیدا نکرد."
    )


def gregorian_to_jalali(gy: int, gm: int, gd: int):
    """
    تبدیل تاریخ میلادی به شمسی بدون نیاز به نصب کتابخانه‌ی اضافه.
    """
    g_day_no = (
        365 * (gy - 1600)
        + (gy - 1600 + 3) // 4
        - (gy - 1600 + 99) // 100
        + (gy - 1600 + 399) // 400
    )

    g_month_days = [31, 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31]

    leap = (
        gy % 4 == 0
        and (gy % 100 != 0 or gy % 400 == 0)
    )

    for month in range(gm - 1):
        g_day_no += g_month_days[month]
        if month == 1 and leap:
            g_day_no += 1

    g_day_no += gd - 1

    j_day_no = g_day_no - 79
    j_np = j_day_no // 12053
    j_day_no %= 12053

    jy = 979 + 33 * j_np + 4 * (j_day_no // 1461)
    j_day_no %= 1461

    if j_day_no >= 366:
        jy += (j_day_no - 1) // 365
        j_day_no = (j_day_no - 1) % 365

    if j_day_no < 186:
        jm = 1 + j_day_no // 31
        jd = 1 + j_day_no % 31
    else:
        jm = 7 + (j_day_no - 186) // 30
        jd = 1 + (j_day_no - 186) % 30

    return jy, jm, jd


def today_jalali_string() -> str:
    today = datetime.now().date()
    jy, jm, jd = gregorian_to_jalali(
        today.year,
        today.month,
        today.day,
    )
    return f"{jy:04d}/{jm:02d}/{jd:02d}"


def normalize_code(value) -> str:
    """
    کد کالا را برای تطبیق آماده می‌کند.
    اگر کد در اکسل عدد باشد، .0 انتهای آن حذف می‌شود.
    اگر کد متن باشد، صفرهای ابتدای آن دستکاری نمی‌شوند.
    """
    if value is None:
        return ""

    if isinstance(value, bool):
        return str(value)

    if isinstance(value, int):
        return str(value)

    if isinstance(value, float):
        if not isfinite(value):
            return ""
        if value.is_integer():
            return str(int(value))
        return format(value, "f").rstrip("0").rstrip(".")

    text = str(value).strip().translate(PERSIAN_ARABIC_DIGITS)
    text = text.replace(",", "").replace("٬", "").strip()

    # فقط 123.0 / 123.00 را به 123 تبدیل می‌کنیم.
    if re.fullmatch(r"\d+\.0+", text):
        return text.split(".", 1)[0]

    return text


def numeric_code_match_key(value) -> str:
    """
    کلید کمکی فقط برای تطبیق کدهای کاملاً عددی.
    خود کد اصلی تغییر نمی‌کند؛ اما 0463 و 463 برای Match یکسان می‌شوند.
    """
    code = normalize_code(value)
    if not code or not re.fullmatch(r"\d+", code):
        return code
    stripped = code.lstrip("0")
    return stripped or "0"


def to_number(value) -> float:
    """
    مقدار موجودی یا ضریب را به عدد تبدیل می‌کند.
    اگر مقدار معتبر نباشد، خطا می‌دهد تا برنامه حدس نزند.
    """
    if value is None:
        raise ValueError("مقدار خالی است")

    if isinstance(value, bool):
        raise ValueError("مقدار True/False عدد موجودی نیست")

    if isinstance(value, (int, float)):
        number = float(value)
        if not isfinite(number):
            raise ValueError("عدد معتبر نیست")
        return number

    text = str(value).strip().translate(PERSIAN_ARABIC_DIGITS)
    text = text.replace(",", "").replace("٬", "").replace(" ", "")

    if text == "":
        raise ValueError("مقدار خالی است")

    try:
        number = float(text)
    except ValueError:
        raise ValueError(f"مقدار عددی نیست: {value!r}")

    if not isfinite(number):
        raise ValueError("عدد معتبر نیست")

    return number


def parse_level(value) -> int:
    """
    سطح موجودی فقط باید 1 یا 2 یا 3 باشد.
    """
    number = to_number(value)

    if not number.is_integer():
        raise ValueError(f"سطح موجودی باید عدد صحیح باشد؛ مقدار فعلی: {value!r}")

    level = int(number)
    if level not in (1, 2, 3):
        raise ValueError(f"سطح موجودی فقط می‌تواند 1 یا 2 یا 3 باشد؛ مقدار فعلی: {value!r}")

    return level


def row_has_any_value(values) -> bool:
    return any(v is not None and str(v).strip() != "" for v in values)


def make_header_positions(header_values) -> dict[str, list[int]]:
    """
    خروجی:
        {
            "کد کالا": [1],
            "موجودی کارتن": [6],
            ...
        }
    شماره ستون‌ها 1-based هستند.
    """
    positions = defaultdict(list)

    for col_idx, value in enumerate(header_values, start=1):
        normalized = normalize_text(value)
        if normalized:
            positions[normalized].append(col_idx)

    return dict(positions)


def make_inventory_header_positions(ws, header_row: int) -> dict[str, list[int]]:
    """
    هدرهای فایل موجودی را با پشتیبانی از Merge برمی‌گرداند.

    مثال فایل زنجان:
      X1:AC1 = «کد کالا»
    در این حالت همه ستون‌های X تا AC به‌عنوان محدوده‌ی کاندید ثبت می‌شوند
    تا بعداً ستون واقعی داده با محتوای ردیف‌ها تشخیص داده شود.
    """
    positions = defaultdict(list)

    # هدرهای عادی / سلول‌های بالای Merge
    for col_idx in range(1, ws.max_column + 1):
        value = ws.cell(header_row, col_idx).value
        normalized = normalize_text(value)
        if not normalized:
            continue

        covered_columns = [col_idx]

        for merged_range in ws.merged_cells.ranges:
            if (
                merged_range.min_row <= header_row <= merged_range.max_row
                and merged_range.min_col == col_idx
                and merged_range.min_row == header_row
            ):
                covered_columns = list(
                    range(merged_range.min_col, merged_range.max_col + 1)
                )
                break

        positions[normalized].extend(covered_columns)

    return {
        key: sorted(set(cols))
        for key, cols in positions.items()
    }


def score_code_column(ws, header_row: int, col_idx: int, known_codes: set[str]):
    """
    امتیاز یک ستون برای اینکه ستون واقعی کد کالا باشد.
    اول تعداد کدهای منطبق با دیکشنری، سپس تعداد سلول‌های غیرخالی.
    """
    scan_end = min(ws.max_row, header_row + 200)
    matches = 0
    nonempty = 0

    for row_idx in range(header_row + 1, scan_end + 1):
        raw = ws.cell(row_idx, col_idx).value
        code = normalize_code(raw)
        if not code:
            continue
        nonempty += 1
        if (
            code in known_codes
            or numeric_code_match_key(code) in known_codes
        ):
            matches += 1

    return matches, nonempty


def resolve_numeric_data_column(ws, header_row: int, candidate_columns: list[int]):
    """
    از داخل محدوده‌ی یک هدر Merge شده، ستونی را پیدا می‌کند که واقعاً
    داده‌ی عددی موجودی در آن قرار دارد.

    خروجی:
      (column_index, None) در حالت موفق
      (None, reason) در حالت مبهم/بدون داده
    """
    if not candidate_columns:
        return None, "ستونی برای این هدر پیدا نشد."

    if len(candidate_columns) == 1:
        return candidate_columns[0], None

    scan_end = min(ws.max_row, header_row + 500)
    scored = []

    for col_idx in candidate_columns:
        numeric_count = 0
        nonempty_count = 0

        for row_idx in range(header_row + 1, scan_end + 1):
            value = ws.cell(row_idx, col_idx).value
            if value is None or str(value).strip() == "":
                continue

            nonempty_count += 1
            try:
                to_decimal(value)
                numeric_count += 1
            except ValueError:
                pass

        scored.append((numeric_count, nonempty_count, col_idx))

    best_numeric = max(item[0] for item in scored)
    best = [item for item in scored if item[0] == best_numeric]

    if best_numeric == 0:
        return None, "در محدوده‌ی هدر Merge شده هیچ ستون عددی معتبری پیدا نشد."

    best_nonempty = max(item[1] for item in best)
    best = [item for item in best if item[1] == best_nonempty]

    if len(best) != 1:
        columns = ", ".join(get_column_letter(item[2]) for item in best)
        return None, f"چند ستون عددی با امتیاز برابر پیدا شد: {columns}"

    return best[0][2], None


def resolve_text_data_column(ws, header_row: int, candidate_columns: list[int]):
    """ستون واقعی متن (مثل نام کالا) را داخل یک هدر Merge شده پیدا می‌کند."""
    if not candidate_columns:
        return None

    if len(candidate_columns) == 1:
        return candidate_columns[0]

    scan_end = min(ws.max_row, header_row + 500)
    scored = []

    for col_idx in candidate_columns:
        nonempty_count = 0
        text_count = 0

        for row_idx in range(header_row + 1, scan_end + 1):
            value = ws.cell(row_idx, col_idx).value
            if value is None or str(value).strip() == "":
                continue
            nonempty_count += 1
            if isinstance(value, str):
                text_count += 1

        scored.append((text_count, nonempty_count, col_idx))

    best_text = max(item[0] for item in scored)
    best = [item for item in scored if item[0] == best_text]

    if best_text == 0:
        best_nonempty = max(item[1] for item in scored)
        best = [item for item in scored if item[1] == best_nonempty]

    if len(best) == 1 and (best[0][0] > 0 or best[0][1] > 0):
        return best[0][2]

    return None


def append_error(
    errors: list[dict],
    *,
    file_name="",
    representative="",
    sheet_name="",
    row_number="",
    code="",
    product_name="",
    error_type="",
    description="",
    dictionary_row="",
    required_column="",
):
    errors.append(
        {
            "نام فایل": file_name,
            "نام نماینده": representative,
            "نام شیت": sheet_name,
            "ردیف فایل": row_number,
            "کد کالا": code,
            "نام کالا در فایل": product_name,
            "نوع خطا": error_type,
            "توضیح": description,
            "ردیف دیکشنری": dictionary_row,
            "ستون مورد نیاز": required_column,
        }
    )


def resolve_dictionary_file() -> Path:
    """
    جدیدترین فایل دیکشنری کنار برنامه را پیدا می‌کند.

    نمونه‌های قابل قبول:
      Dictionary_v2_matched_updated.xlsx
      Dictionary_v2_matched_updated(2).xlsx
      Dictionary_v2_matched_updated(7).xlsx

    اگر چند نسخه وجود داشته باشد، فایلی که آخرین بار تغییر کرده انتخاب می‌شود
    و نامش در ترمینال چاپ می‌شود.
    """
    candidates = [
        path
        for path in BASE_DIR.glob(f"{DICTIONARY_BASENAME}*.xlsx")
        if path.is_file()
        and not path.name.startswith("~$")
        and path.resolve() != OUTPUT_FILE.resolve()
    ]

    if not candidates:
        raise FileNotFoundError(
            "هیچ فایل دیکشنری پیدا نشد. "
            f"نام فایل باید با «{DICTIONARY_BASENAME}» شروع شود."
        )

    candidates.sort(
        key=lambda p: (
            p.stat().st_mtime_ns,
            p.name,
        ),
        reverse=True,
    )

    return candidates[0]


# ============================================================
# خواندن دیکشنری
# ============================================================

def find_dictionary_table(workbook):
    required_headers = {
        normalize_text(D_REP),
        normalize_text(D_REP_CODE),
        normalize_text(D_COMPANY_CODE),
        normalize_text(D_COMPANY_NAME),
        normalize_text(D_OWN_CODE),
        normalize_text(D_SHORT_NAME),
        normalize_text(D_BRAND),
        normalize_text(D_PRODUCT_GROUP),
        normalize_text(D_FLAVOR),
        normalize_text(D_WEIGHT),
        normalize_text(D_LEVEL),
        normalize_text(D_INPUT_COLUMN),
        normalize_text(D_MIDDLE_PER_MASTER),
        normalize_text(D_UNITS_PER_CARTON),
    }

    candidates = []

    for ws in workbook.worksheets:
        max_scan = min(HEADER_SCAN_LIMIT, ws.max_row)

        for row_idx in range(1, max_scan + 1):
            values = [ws.cell(row_idx, col).value for col in range(1, ws.max_column + 1)]
            normalized = {normalize_text(v) for v in values if normalize_text(v)}

            if required_headers.issubset(normalized):
                candidates.append((ws, row_idx))

    if not candidates:
        raise RuntimeError(
            "هدر دیکشنری پیدا نشد. این ستون‌ها باید وجود داشته باشند:\n"
            + "\n".join(sorted(required_headers))
        )

    if len(candidates) > 1:
        locations = ", ".join(f"{ws.title}!{row}" for ws, row in candidates)
        raise RuntimeError(
            "بیش از یک جدول شبیه دیکشنری پیدا شد و برنامه نباید حدس بزند:\n"
            f"{locations}"
        )

    return candidates[0]


def load_dictionary(path: Path):
    if not path.exists():
        raise FileNotFoundError(f"فایل دیکشنری پیدا نشد:\n{path}")

    wb = load_workbook(BytesIO(path.read_bytes()), read_only=False, data_only=True)
    ws, header_row = find_dictionary_table(wb)

    header_values = [
        ws.cell(header_row, col).value
        for col in range(1, ws.max_column + 1)
    ]
    header_positions = make_header_positions(header_values)

    def get_unique_col(header_name: str) -> int:
        normalized = normalize_text(header_name)
        positions = header_positions.get(normalized, [])

        if not positions:
            raise RuntimeError(f"ستون «{header_name}» در دیکشنری پیدا نشد.")

        if len(positions) > 1:
            raise RuntimeError(f"ستون «{header_name}» در دیکشنری تکراری است.")

        return positions[0]

    cols = {
        D_REP: get_unique_col(D_REP),
        D_REP_CODE: get_unique_col(D_REP_CODE),
        D_COMPANY_CODE: get_unique_col(D_COMPANY_CODE),
        D_COMPANY_NAME: get_unique_col(D_COMPANY_NAME),
        D_OWN_CODE: get_unique_col(D_OWN_CODE),
        D_SHORT_NAME: get_unique_col(D_SHORT_NAME),
        D_BRAND: get_unique_col(D_BRAND),
        D_PRODUCT_GROUP: get_unique_col(D_PRODUCT_GROUP),
        D_FLAVOR: get_unique_col(D_FLAVOR),
        D_WEIGHT: get_unique_col(D_WEIGHT),
        D_LEVEL: get_unique_col(D_LEVEL),
        D_INPUT_COLUMN: get_unique_col(D_INPUT_COLUMN),
        D_MIDDLE_PER_MASTER: get_unique_col(D_MIDDLE_PER_MASTER),
        D_UNITS_PER_CARTON: get_unique_col(D_UNITS_PER_CARTON),
    }

    rep_product_positions = header_positions.get(
        normalize_text(D_REP_PRODUCT),
        [],
    )
    cols[D_REP_PRODUCT] = (
        rep_product_positions[0]
        if len(rep_product_positions) == 1
        else None
    )

    entries_by_rep_code = defaultdict(lambda: defaultdict(list))
    representatives = set()
    representative_aliases = defaultdict(set)
    representative_info_sets = defaultdict(set)
    all_entries = []

    for row_idx in range(header_row + 1, ws.max_row + 1):
        representative_raw = ws.cell(row_idx, cols[D_REP]).value
        code_raw = ws.cell(row_idx, cols[D_REP_CODE]).value

        if representative_raw is None and code_raw is None:
            continue

        representative = (
            ""
            if representative_raw is None
            else str(representative_raw).strip()
        )
        code = normalize_code(code_raw)
        own_code = normalize_code(
            ws.cell(row_idx, cols[D_OWN_CODE]).value
        )

        company_code = normalize_code(
            ws.cell(row_idx, cols[D_COMPANY_CODE]).value
        )
        company_name = normalize_text(
            ws.cell(row_idx, cols[D_COMPANY_NAME]).value
        )

        entry = {
            "dictionary_row": row_idx,
            "representative": representative,
            "company_code": company_code,
            "company_name": company_name,
            "code": code,
            "own_code": own_code,
            "representative_product": (
                ws.cell(row_idx, cols[D_REP_PRODUCT]).value
                if cols[D_REP_PRODUCT]
                else None
            ),
            "short_name": normalize_text(
                ws.cell(row_idx, cols[D_SHORT_NAME]).value
            ),
            "brand": normalize_text(
                ws.cell(row_idx, cols[D_BRAND]).value
            ),
            "product_group": normalize_text(
                ws.cell(row_idx, cols[D_PRODUCT_GROUP]).value
            ),
            "flavor": normalize_text(
                ws.cell(row_idx, cols[D_FLAVOR]).value
            ),
            "weight": ws.cell(row_idx, cols[D_WEIGHT]).value,
            "level_raw": ws.cell(row_idx, cols[D_LEVEL]).value,
            "input_column": normalize_text(
                ws.cell(row_idx, cols[D_INPUT_COLUMN]).value
            ),
            "middle_per_master_raw": ws.cell(
                row_idx,
                cols[D_MIDDLE_PER_MASTER],
            ).value,
            "units_per_carton_raw": ws.cell(
                row_idx,
                cols[D_UNITS_PER_CARTON],
            ).value,
        }

        all_entries.append(entry)

        if representative:
            representatives.add(representative)
            representative_aliases[
                representative_match_key(representative)
            ].add(representative)

            if company_code or company_name:
                representative_info_sets[representative].add(
                    (company_code, company_name)
                )

        if representative and code:
            entries_by_rep_code[representative][code].append(entry)

    # اطلاعات شرکت نماینده:
    # نام کامل اگر یکتا باشد استفاده می‌شود.
    # کد شرکت اگر یکتا نباشد خالی می‌ماند تا برنامه حدس نزند.
    representative_info = {}

    for representative in representatives:
        infos = representative_info_sets.get(representative, set())

        company_codes = {
            code
            for code, name in infos
            if code
        }
        company_names = {
            name
            for code, name in infos
            if name
        }

        company_code = (
            next(iter(company_codes))
            if len(company_codes) == 1
            else ""
        )

        company_name = (
            next(iter(company_names))
            if len(company_names) == 1
            else representative
        )

        representative_info[representative] = {
            "company_code": company_code,
            "company_name": company_name,
            "company_code_ambiguous": len(company_codes) > 1,
            "company_name_ambiguous": len(company_names) > 1,
        }

    # --------------------------------------------------------
    # تعیین نام نمایشی طعم.
    # فقط وقتی وزن‌های متفاوت واقعاً به کدهای خودمان متفاوت مربوط باشند،
    # وزن در عنوان می‌آید. بنابراین خطای وزن 250 برای نی‌شیر باعث
    # ساخت ستون جدا نمی‌شود چون همان کد کالای خودمان است.
    # --------------------------------------------------------
    metadata_by_identity = defaultdict(
        lambda: {
            "weights": set(),
            "own_codes": set(),
        }
    )

    for entry in all_entries:
        brand = normalize_text(entry["brand"])
        group = normalize_text(entry["product_group"])
        section, base_flavor = section_and_base_flavor(entry)
        weight = normalize_weight(entry["weight"])
        own_code = normalize_code(entry["own_code"])

        if not brand or not group or not base_flavor:
            continue

        identity = (
            brand,
            group,
            section,
            base_flavor,
        )

        if weight:
            metadata_by_identity[identity]["weights"].add(weight)

        if own_code:
            metadata_by_identity[identity]["own_codes"].add(own_code)

    output_key_set = set()

    for entry in all_entries:
        brand = normalize_text(entry["brand"])
        group = normalize_text(entry["product_group"])
        section, base_flavor = section_and_base_flavor(entry)
        weight = normalize_weight(entry["weight"])

        if not brand or not group or not base_flavor:
            continue

        identity = (
            brand,
            group,
            section,
            base_flavor,
        )
        meta = metadata_by_identity[identity]

        add_weight = (
            base_flavor != "تک طعم"
            and group != "نی شیر"
            and weight
            and len(meta["weights"]) > 1
            and len(meta["own_codes"]) > 1
        )

        if add_weight:
            # نگارش فارسی: «180 گرمی»، نه «گرم 180»
            display_flavor = f"{base_flavor} {weight} گرمی"
        else:
            display_flavor = base_flavor

        entry["display_flavor"] = display_flavor

        output_key_set.add(
            (
                brand,
                group,
                section,
                display_flavor,
            )
        )

    # برند -> گروه کالا -> سکشن -> طعم
    # هر مورد «جور» آخر همان سکشن می‌رود.
    output_keys = sorted(
        output_key_set,
        key=lambda key: (
            persian_sort_key(key[0]),
            persian_sort_key(key[1]),
            persian_sort_key(key[2]),
            flavor_sort_key(key[3]),
        ),
    )

    wb.close()

    return {
        "entries_by_rep_code": entries_by_rep_code,
        "output_keys": output_keys,
        "representatives": representatives,
        "representative_aliases": representative_aliases,
        "representative_info": representative_info,
    }


# ============================================================
# پیدا کردن جدول موجودی نماینده
# ============================================================

def find_inventory_table(workbook, representative_entries):
    """
    جدول موجودی را پیدا می‌کند.

    حالت اول:
      ستون کد هدر صریح دارد، مثل «کد کالا»، «کدکالا» یا «کد».

    حالت دوم:
      هدر ستون کد خالی است (مثل فایل تبریز رونق).
      در این حالت برنامه ستون کد را با تطبیق واقعی مقادیر همان ستون
      با کدهای موجود در دیکشنری همان نماینده پیدا می‌کند.
      اگر بیش از یک ستون به یک اندازه محتمل باشد، برنامه حدس نمی‌زند.
    """

    expected_columns = {
        normalize_text(entry["input_column"])
        for code_entries in representative_entries.values()
        for entry in code_entries
        if normalize_text(entry["input_column"])
    }

    known_rep_codes = set()
    for code in representative_entries.keys():
        normalized_code = normalize_code(code)
        if not normalized_code:
            continue
        known_rep_codes.add(normalized_code)
        known_rep_codes.add(numeric_code_match_key(normalized_code))

    code_headers_normalized = {
        normalize_text(x)
        for x in CODE_COLUMN_CANDIDATES
    }

    explicit_candidates = []

    # --------------------------------------------------------
    # مرحله 1: پیدا کردن هدر صریح کد
    # --------------------------------------------------------
    for ws in workbook.worksheets:
        max_scan = min(HEADER_SCAN_LIMIT, ws.max_row)

        for row_idx in range(1, max_scan + 1):
            header_values = [
                ws.cell(row_idx, col).value
                for col in range(1, ws.max_column + 1)
            ]
            header_positions = make_inventory_header_positions(ws, row_idx)

            code_column_indexes = []

            for candidate in code_headers_normalized:
                code_column_indexes.extend(
                    header_positions.get(candidate, [])
                )

            code_column_indexes = sorted(
                set(code_column_indexes)
            )

            if not code_column_indexes:
                continue

            # اگر هدر کد Merge شده باشد (مثل X1:AC1 در فایل زنجان)،
            # ستون واقعی را با تطبیق محتوای آن با کدهای دیکشنری پیدا می‌کنیم.
            scored_code_columns = [
                (
                    *score_code_column(
                        ws,
                        row_idx,
                        col_idx,
                        known_rep_codes,
                    ),
                    col_idx,
                )
                for col_idx in code_column_indexes
            ]

            best_code_matches = max(item[0] for item in scored_code_columns)
            best_code_cols = [
                item for item in scored_code_columns
                if item[0] == best_code_matches
            ]

            if len(best_code_cols) > 1:
                best_nonempty = max(item[1] for item in best_code_cols)
                best_code_cols = [
                    item for item in best_code_cols
                    if item[1] == best_nonempty
                ]

            if len(best_code_cols) != 1:
                continue

            selected_code_col = best_code_cols[0][2]

            # اگر هدر Merge بوده، حداقل یک تطبیق واقعی لازم است تا
            # ستون اشتباهی از محدوده‌ی Merge انتخاب نشود.
            if len(code_column_indexes) > 1 and best_code_matches == 0:
                continue

            matched_expected_columns = sum(
                1
                for col_name in expected_columns
                if col_name in header_positions
            )

            explicit_candidates.append(
                {
                    "sheet": ws,
                    "header_row": row_idx,
                    "header_positions": header_positions,
                    "code_col": selected_code_col,
                    "score": matched_expected_columns,
                    "code_match_count": best_code_matches,
                    "detection_mode": "header",
                }
            )

    if explicit_candidates:
        best_score = max(
            item["score"]
            for item in explicit_candidates
        )
        best = [
            item
            for item in explicit_candidates
            if item["score"] == best_score
        ]

        if len(best) > 1:
            locations = ", ".join(
                f"{item['sheet'].title}!ردیف {item['header_row']}"
                for item in best
            )
            raise RuntimeError(
                "چند جدول محتمل با شرایط برابر پیدا شد و برنامه نباید حدس بزند:\n"
                f"{locations}"
            )

        return best[0]

    # --------------------------------------------------------
    # مرحله 2: هدر کد نداریم.
    # ستون کد را با خودِ مقادیر دیکشنری همان نماینده پیدا می‌کنیم.
    # --------------------------------------------------------
    inferred_candidates = []

    for ws in workbook.worksheets:
        max_scan = min(HEADER_SCAN_LIMIT, ws.max_row)

        for header_row in range(1, max_scan + 1):
            header_values = [
                ws.cell(header_row, col).value
                for col in range(1, ws.max_column + 1)
            ]
            header_positions = make_inventory_header_positions(
                ws,
                header_row,
            )

            matched_expected_columns = sum(
                1
                for col_name in expected_columns
                if col_name in header_positions
            )

            # برای حالت بدون هدر کد، حداقل یکی از ستون‌های
            # مورد انتظار دیکشنری باید در این ردیف هدر دیده شود.
            if expected_columns and matched_expected_columns == 0:
                continue

            scan_end = min(
                ws.max_row,
                header_row + 100,
            )

            column_scores = []

            for col_idx in range(1, ws.max_column + 1):
                matches = 0
                nonempty = 0

                for row_idx in range(
                    header_row + 1,
                    scan_end + 1,
                ):
                    raw_value = ws.cell(
                        row_idx,
                        col_idx,
                    ).value

                    if raw_value is None:
                        continue

                    code_value = normalize_code(
                        raw_value
                    )

                    if not code_value:
                        continue

                    nonempty += 1

                    if code_value in known_rep_codes:
                        matches += 1

                if matches > 0:
                    column_scores.append(
                        {
                            "col": col_idx,
                            "matches": matches,
                            "nonempty": nonempty,
                        }
                    )

            if not column_scores:
                continue

            max_matches = max(
                item["matches"]
                for item in column_scores
            )

            best_columns = [
                item
                for item in column_scores
                if item["matches"] == max_matches
            ]

            # اگر دو ستون به یک اندازه با دیکشنری جور باشند،
            # این ردیف هدر را مبهم در نظر می‌گیریم.
            if len(best_columns) != 1:
                continue

            best_col = best_columns[0]

            # حداقل 2 کد واقعی باید با دیکشنری تطبیق داشته باشند
            # تا یک عدد اتفاقی به‌عنوان ستون کد انتخاب نشود.
            if best_col["matches"] < 2:
                continue

            inferred_candidates.append(
                {
                    "sheet": ws,
                    "header_row": header_row,
                    "header_positions": header_positions,
                    "code_col": best_col["col"],
                    "score": matched_expected_columns,
                    "code_match_count": best_col["matches"],
                    "detection_mode": "dictionary_match",
                }
            )

    if not inferred_candidates:
        raise RuntimeError(
            "هیچ جدول موجودی با ستون کد معتبر پیدا نشد. "
            "نه هدر کد شناخته‌شده وجود داشت و نه ستونی پیدا شد "
            "که مقادیرش به‌طور یکتا با کدهای دیکشنری این نماینده تطبیق داشته باشد. "
            f"عنوان‌های کد مورد قبول فعلی: {sorted(CODE_COLUMN_CANDIDATES)}"
        )

    # اول تعداد تطبیق واقعی کدها، بعد تعداد ستون‌های مورد انتظار هدر
    best_code_match = max(
        item["code_match_count"]
        for item in inferred_candidates
    )

    best = [
        item
        for item in inferred_candidates
        if item["code_match_count"] == best_code_match
    ]

    if len(best) > 1:
        best_score = max(
            item["score"]
            for item in best
        )
        best = [
            item
            for item in best
            if item["score"] == best_score
        ]

    if len(best) > 1:
        locations = ", ".join(
            (
                f"{item['sheet'].title}!ردیف {item['header_row']} "
                f"(ستون {get_column_letter(item['code_col'])}, "
                f"{item['code_match_count']} تطبیق)"
            )
            for item in best
        )
        raise RuntimeError(
            "چند ستون کد محتمل با شرایط برابر پیدا شد و برنامه نباید حدس بزند:\n"
            f"{locations}"
        )

    return best[0]


def find_optional_product_name_column(
    ws,
    header_row: int,
    header_positions: dict[str, list[int]],
):
    """
    فقط برای نمایش نام کالا در شیت خطاها.
    از هدرهای Merge شده هم ستون واقعی نام کالا را پیدا می‌کند.
    """
    for candidate in ("نام کالا", "کالا"):
        positions = header_positions.get(normalize_text(candidate), [])
        resolved = resolve_text_data_column(
            ws,
            header_row,
            positions,
        )
        if resolved is not None:
            return resolved
    return None


# ============================================================
# تبدیل هر فایل نماینده
# ============================================================

def process_inventory_file(
    path: Path,
    dictionary_data: dict,
    errors: list[dict],
):
    input_representative_name = path.stem

    try:
        representative = resolve_representative(
            input_representative_name,
            dictionary_data,
        )
    except ValueError as exc:
        append_error(
            errors,
            file_name=path.name,
            representative=input_representative_name,
            error_type="نام نماینده پیدا نشد",
            description=str(exc),
        )
        return None

    representative_entries = dictionary_data["entries_by_rep_code"][representative]

    # Alias امن برای تفاوت صفرهای ابتدایی کدهای کاملاً عددی.
    # مثال: فایل نماینده 0463 ولی دیکشنری 463.
    numeric_code_aliases = defaultdict(set)
    for dictionary_code in representative_entries.keys():
        alias = numeric_code_match_key(dictionary_code)
        if alias:
            numeric_code_aliases[alias].add(dictionary_code)

    rep_info = dictionary_data["representative_info"].get(
        representative,
        {
            "company_code": "",
            "company_name": representative,
            "company_code_ambiguous": False,
            "company_name_ambiguous": False,
        },
    )

    if rep_info.get("company_code_ambiguous"):
        append_error(
            errors,
            file_name=path.name,
            representative=representative,
            error_type="کد شرکت مبهم است",
            description=(
                "برای این نماینده بیش از یک «کد شرکت» در دیکشنری ثبت شده است. "
                "برای جلوگیری از حدس، ستون کد شرکت در خروجی خالی گذاشته شد."
            ),
        )

    if rep_info.get("company_name_ambiguous"):
        append_error(
            errors,
            file_name=path.name,
            representative=representative,
            error_type="نام کامل نماینده مبهم است",
            description=(
                "برای این نماینده بیش از یک «نام شرکت» در دیکشنری ثبت شده است. "
                "برای جلوگیری از حدس، نام کوتاه نماینده در خروجی استفاده شد."
            ),
        )

    try:
        wb = load_workbook(BytesIO(path.read_bytes()), read_only=False, data_only=True)
    except Exception as exc:
        append_error(
            errors,
            file_name=path.name,
            representative=representative,
            error_type="خطا در باز کردن فایل",
            description=str(exc),
        )
        return None

    try:
        table = find_inventory_table(wb, representative_entries)
    except Exception as exc:
        append_error(
            errors,
            file_name=path.name,
            representative=representative,
            error_type="ساختار فایل قابل تشخیص نیست",
            description=str(exc),
        )
        wb.close()
        return None

    ws = table["sheet"]
    header_row = table["header_row"]
    header_positions = table["header_positions"]
    code_col = table["code_col"]
    product_name_col = find_optional_product_name_column(
        ws,
        header_row,
        header_positions,
    )

    # ستون‌های عددی موردنیاز را یک بار Resolve می‌کنیم تا در فایل‌های
    # دارای هدر Merge شده، ستون واقعی داده (مثلاً M زیر L:N) پیدا شود.
    resolved_input_columns = {}

    records = []

    for row_idx in range(header_row + 1, ws.max_row + 1):
        row_values = [
            ws.cell(row_idx, col).value
            for col in range(1, ws.max_column + 1)
        ]

        if not row_has_any_value(row_values):
            continue

        raw_code = ws.cell(row_idx, code_col).value
        code = normalize_code(raw_code)

        product_name = (
            ws.cell(row_idx, product_name_col).value
            if product_name_col
            else ""
        )

        if not code:
            # ردیف جمع/فوتر که نام کالا هم ندارد، داده کالایی نیست.
            if product_name is None or str(product_name).strip() == "":
                continue

            append_error(
                errors,
                file_name=path.name,
                representative=representative,
                sheet_name=ws.title,
                row_number=row_idx,
                product_name=product_name,
                error_type="کد کالا خالی است",
                description="ردیف نام کالا دارد ولی ستون کد کالا خالی است.",
            )
            continue

        records.append(
            {
                "row": row_idx,
                "code": code,
                "product_name": product_name,
            }
        )

    # همه ردیف‌ها—even اگر خودِ کد نماینده تکراری باشد—پردازش می‌شوند.
    # در نهایت همه چیز بر اساس «کد کالای خودمان» تجمیع می‌شود.
    # بنابراین:
    # 1) یک کد نماینده که چند بار در فایل آمده باشد، جمع می‌شود.
    # 2) چند کد متفاوت نماینده که به یک کد خودمان برسند نیز با هم جمع می‌شوند.
    exact_by_own_code = defaultdict(list)

    for record in records:
        row_idx = record["row"]
        code = record["code"]
        product_name = record["product_name"]

        dictionary_matches = representative_entries.get(code, [])

        if not dictionary_matches:
            alias = numeric_code_match_key(code)
            alias_candidates = numeric_code_aliases.get(alias, set())

            if len(alias_candidates) == 1:
                matched_dictionary_code = next(iter(alias_candidates))
                dictionary_matches = representative_entries.get(
                    matched_dictionary_code,
                    [],
                )
            elif len(alias_candidates) > 1:
                append_error(
                    errors,
                    file_name=path.name,
                    representative=representative,
                    sheet_name=ws.title,
                    row_number=row_idx,
                    code=code,
                    product_name=product_name,
                    error_type="کد با صفر ابتدایی مبهم است",
                    description=(
                        "بعد از حذف صفرهای ابتدایی، این کد به بیش از یک کد "
                        "در دیکشنری می‌خورد؛ برنامه برای جلوگیری از حدس محاسبه نکرد."
                    ),
                )
                continue

        if not dictionary_matches:
            append_error(
                errors,
                file_name=path.name,
                representative=representative,
                sheet_name=ws.title,
                row_number=row_idx,
                code=code,
                product_name=product_name,
                error_type="کد در دیکشنری پیدا نشد",
                description=(
                    "برای این نماینده و این کد کالا هیچ ردیفی "
                    "در دیکشنری وجود ندارد."
                ),
            )
            continue

        if len(dictionary_matches) > 1:
            # یک کد نماینده ممکن است در دیکشنری چند تعریف جایگزین داشته باشد؛
            # مثلاً «عدد / سطح 3» و «کارتن / سطح 1».
            #
            # قاعده‌ی قطعی:
            # منبع جزئی‌تر ارجح است: سطح 3 > سطح 2 > سطح 1.
            # بنابراین اگر فقط یک تعریف در بالاترین سطح معتبر باشد،
            # همان انتخاب می‌شود و تعاریف سطح پایین‌تر نادیده گرفته می‌شوند.
            #
            # این دقیقاً برای بندرعباس باعث می‌شود «عدد» مرجع باشد.
            valid_candidates = []

            for candidate in dictionary_matches:
                try:
                    level_candidate = parse_level(candidate["level_raw"])
                except ValueError:
                    continue

                input_name = normalize_text(candidate["input_column"])

                if input_name not in resolved_input_columns:
                    positions = header_positions.get(input_name, [])
                    resolved_col, _ = resolve_numeric_data_column(
                        ws,
                        header_row,
                        positions,
                    )
                    resolved_input_columns[input_name] = resolved_col

                resolved_col = resolved_input_columns.get(input_name)

                if resolved_col is None:
                    continue

                try:
                    quantity_candidate = to_decimal(
                        ws.cell(row_idx, resolved_col).value
                    )
                except ValueError:
                    continue

                if quantity_candidate < 0:
                    continue

                if level_candidate == 1:
                    converted_candidate = quantity_candidate

                elif level_candidate == 2:
                    try:
                        divisor_candidate = to_decimal(
                            candidate["middle_per_master_raw"]
                        )
                    except ValueError:
                        continue

                    if divisor_candidate <= 0:
                        continue

                    converted_candidate = (
                        quantity_candidate / divisor_candidate
                    )

                else:  # level 3
                    try:
                        divisor_candidate = to_decimal(
                            candidate["units_per_carton_raw"]
                        )
                    except ValueError:
                        continue

                    if divisor_candidate <= 0:
                        continue

                    converted_candidate = (
                        quantity_candidate / divisor_candidate
                    )

                valid_candidates.append(
                    {
                        "entry": candidate,
                        "level": level_candidate,
                        "converted": converted_candidate,
                    }
                )

            if not valid_candidates:
                dict_rows = ", ".join(
                    str(x["dictionary_row"])
                    for x in dictionary_matches
                )

                append_error(
                    errors,
                    file_name=path.name,
                    representative=representative,
                    sheet_name=ws.title,
                    row_number=row_idx,
                    code=code,
                    product_name=product_name,
                    error_type="تعریف تکراری بدون گزینه معتبر",
                    description=(
                        "چند تعریف در دیکشنری وجود دارد اما هیچ‌کدام "
                        "با ساختار و مقادیر فایل ورودی قابل استفاده نیست."
                    ),
                    dictionary_row=dict_rows,
                )
                continue

            highest_level = max(
                item["level"]
                for item in valid_candidates
            )

            best = [
                item
                for item in valid_candidates
                if item["level"] == highest_level
            ]

            if len(best) == 1:
                dictionary_matches = [best[0]["entry"]]

            else:
                # اگر در بالاترین سطح چند ردیف کاملاً یکسان باشند،
                # فقط یک بار استفاده می‌کنیم.
                signatures = {
                    (
                        item["entry"]["own_code"],
                        normalize_text(item["entry"]["input_column"]),
                        item["level"],
                        normalize_text(item["entry"]["brand"]),
                        normalize_text(item["entry"]["product_group"]),
                        normalize_text(item["entry"].get("flavor")),
                    )
                    for item in best
                }

                if len(signatures) == 1:
                    dictionary_matches = [best[0]["entry"]]
                else:
                    dict_rows = ", ".join(
                        str(x["dictionary_row"])
                        for x in dictionary_matches
                    )

                    append_error(
                        errors,
                        file_name=path.name,
                        representative=representative,
                        sheet_name=ws.title,
                        row_number=row_idx,
                        code=code,
                        product_name=product_name,
                        error_type="تعریف تکراری مبهم در بالاترین سطح",
                        description=(
                            "در بالاترین سطح موجودی بیش از یک تعریف متفاوت "
                            "وجود دارد؛ برنامه برای جلوگیری از حدس محاسبه نکرد."
                        ),
                        dictionary_row=dict_rows,
                    )
                    continue

        entry = dictionary_matches[0]
        dict_row = entry["dictionary_row"]

        if not entry["own_code"]:
            append_error(
                errors,
                file_name=path.name,
                representative=representative,
                sheet_name=ws.title,
                row_number=row_idx,
                code=code,
                product_name=product_name,
                error_type="کد کالای خودمان خالی است",
                description=(
                    "ستون «کد کالای خودمان» برای این ردیف دیکشنری خالی است."
                ),
                dictionary_row=dict_row,
            )
            continue

        if not entry["brand"]:
            append_error(
                errors,
                file_name=path.name,
                representative=representative,
                sheet_name=ws.title,
                row_number=row_idx,
                code=code,
                product_name=product_name,
                error_type="برند خالی است",
                description="ستون «برند» در دیکشنری برای این کالا خالی است.",
                dictionary_row=dict_row,
            )
            continue

        if not entry["product_group"]:
            append_error(
                errors,
                file_name=path.name,
                representative=representative,
                sheet_name=ws.title,
                row_number=row_idx,
                code=code,
                product_name=product_name,
                error_type="گروه کالا خالی است",
                description="ستون «گروه کالا» در دیکشنری برای این کالا خالی است.",
                dictionary_row=dict_row,
            )
            continue

        if not entry["input_column"]:
            append_error(
                errors,
                file_name=path.name,
                representative=representative,
                sheet_name=ws.title,
                row_number=row_idx,
                code=code,
                product_name=product_name,
                error_type="ستون مورد نیاز خالی است",
                description=(
                    "ستون «ستون مورد نیاز» در دیکشنری برای این کالا خالی است."
                ),
                dictionary_row=dict_row,
            )
            continue

        try:
            level = parse_level(entry["level_raw"])
        except ValueError as exc:
            append_error(
                errors,
                file_name=path.name,
                representative=representative,
                sheet_name=ws.title,
                row_number=row_idx,
                code=code,
                product_name=product_name,
                error_type="سطح موجودی نامعتبر است",
                description=str(exc),
                dictionary_row=dict_row,
                required_column=entry["input_column"],
            )
            continue

        input_column_norm = normalize_text(entry["input_column"])
        input_positions = header_positions.get(input_column_norm, [])

        if not input_positions:
            append_error(
                errors,
                file_name=path.name,
                representative=representative,
                sheet_name=ws.title,
                row_number=row_idx,
                code=code,
                product_name=product_name,
                error_type="ستون موجودی پیدا نشد",
                description=(
                    f"طبق دیکشنری باید ستون «{entry['input_column']}» "
                    "در فایل نماینده وجود داشته باشد."
                ),
                dictionary_row=dict_row,
                required_column=entry["input_column"],
            )
            continue

        if input_column_norm not in resolved_input_columns:
            resolved_col, resolve_error = resolve_numeric_data_column(
                ws,
                header_row,
                input_positions,
            )
            resolved_input_columns[input_column_norm] = resolved_col
        else:
            resolved_col = resolved_input_columns[input_column_norm]
            resolve_error = None

        if resolved_col is None:
            append_error(
                errors,
                file_name=path.name,
                representative=representative,
                sheet_name=ws.title,
                row_number=row_idx,
                code=code,
                product_name=product_name,
                error_type="ستون موجودی Merge شده مبهم است",
                description=(
                    f"عنوان ستون «{entry['input_column']}» روی چند ستون قرار دارد "
                    f"اما ستون واقعی داده با اطمینان تشخیص داده نشد. {resolve_error or ''}"
                ),
                dictionary_row=dict_row,
                required_column=entry["input_column"],
            )
            continue

        raw_quantity = ws.cell(
            row_idx,
            resolved_col,
        ).value

        try:
            quantity = to_decimal(raw_quantity)
        except ValueError as exc:
            append_error(
                errors,
                file_name=path.name,
                representative=representative,
                sheet_name=ws.title,
                row_number=row_idx,
                code=code,
                product_name=product_name,
                error_type="موجودی نامعتبر است",
                description=str(exc),
                dictionary_row=dict_row,
                required_column=entry["input_column"],
            )
            continue

        if quantity < 0:
            append_error(
                errors,
                file_name=path.name,
                representative=representative,
                sheet_name=ws.title,
                row_number=row_idx,
                code=code,
                product_name=product_name,
                error_type="موجودی منفی است",
                description="برای جلوگیری از حدس، موجودی منفی محاسبه نشد.",
                dictionary_row=dict_row,
                required_column=entry["input_column"],
            )
            continue

        # تبدیل دقیق به کارتن مادر؛ هنوز گرد نمی‌کنیم.
        if level == 1:
            converted_exact = Fraction(quantity)

        elif level == 2:
            try:
                divisor = to_decimal(entry["middle_per_master_raw"])
            except ValueError as exc:
                append_error(
                    errors,
                    file_name=path.name,
                    representative=representative,
                    sheet_name=ws.title,
                    row_number=row_idx,
                    code=code,
                    product_name=product_name,
                    error_type="ضریب سطح 2 نامعتبر است",
                    description=(
                        "ستون «تعداد میدل در کارتن مادر» معتبر نیست: "
                        + str(exc)
                    ),
                    dictionary_row=dict_row,
                    required_column=entry["input_column"],
                )
                continue

            if divisor <= 0:
                append_error(
                    errors,
                    file_name=path.name,
                    representative=representative,
                    sheet_name=ws.title,
                    row_number=row_idx,
                    code=code,
                    product_name=product_name,
                    error_type="ضریب سطح 2 نامعتبر است",
                    description=(
                        "تعداد میدل در کارتن مادر باید بزرگ‌تر از صفر باشد."
                    ),
                    dictionary_row=dict_row,
                    required_column=entry["input_column"],
                )
                continue

            converted_exact = Fraction(quantity) / Fraction(divisor)

        else:  # level == 3
            try:
                divisor = to_decimal(entry["units_per_carton_raw"])
            except ValueError as exc:
                append_error(
                    errors,
                    file_name=path.name,
                    representative=representative,
                    sheet_name=ws.title,
                    row_number=row_idx,
                    code=code,
                    product_name=product_name,
                    error_type="ضریب سطح 3 نامعتبر است",
                    description=(
                        "ستون «تعداد در کارتن» معتبر نیست: "
                        + str(exc)
                    ),
                    dictionary_row=dict_row,
                    required_column=entry["input_column"],
                )
                continue

            if divisor <= 0:
                append_error(
                    errors,
                    file_name=path.name,
                    representative=representative,
                    sheet_name=ws.title,
                    row_number=row_idx,
                    code=code,
                    product_name=product_name,
                    error_type="ضریب سطح 3 نامعتبر است",
                    description="تعداد در کارتن باید بزرگ‌تر از صفر باشد.",
                    dictionary_row=dict_row,
                    required_column=entry["input_column"],
                )
                continue

            converted_exact = Fraction(quantity) / Fraction(divisor)

        exact_by_own_code[entry["own_code"]].append(
            {
                "converted_exact": converted_exact,
                "entry": entry,
                "source_row": row_idx,
                "rep_code": code,
                "product_name": product_name,
            }
        )

    totals_by_output_key = defaultdict(int)
    consolidated_own_codes = 0

    for own_code, items in exact_by_own_code.items():
        output_keys = {
            output_key_from_entry(item["entry"])
            for item in items
        }

        # یک کد خودمان باید فقط به یک جای هرم خروجی برسد.
        if len(output_keys) != 1:
            source_rows = ", ".join(
                str(item["source_row"])
                for item in items
            )
            rep_codes = ", ".join(
                sorted(
                    {item["rep_code"] for item in items},
                    key=persian_sort_key,
                )
            )
            append_error(
                errors,
                file_name=path.name,
                representative=representative,
                sheet_name=ws.title,
                row_number=source_rows,
                code=rep_codes,
                error_type="اطلاعات کد خودمان ناسازگار است",
                description=(
                    f"کد کالای خودمان {own_code} از چند کد نماینده "
                    "به برند/گروه/طعم متفاوت رسیده است؛ محاسبه نشد."
                ),
            )
            continue

        if len(items) > 1:
            consolidated_own_codes += 1

        exact_total = sum(
            (item["converted_exact"] for item in items),
            Fraction(0),
        )

        # طبق قاعده تأییدشده: در پایان رو به پایین گرد می‌شود.
        final_quantity = int(
            exact_total.numerator // exact_total.denominator
        )

        output_key = next(iter(output_keys))
        totals_by_output_key[output_key] += final_quantity

    wb.close()

    return {
        "representative": representative,
        "company_code": rep_info["company_code"],
        "company_name": rep_info["company_name"],
        "totals": dict(totals_by_output_key),
        "consolidated_own_codes": consolidated_own_codes,
    }


# ============================================================
# ساخت فایل خروجی
# ============================================================

def style_output_sheet(ws, max_col: int, max_row: int):
    ws.sheet_view.rightToLeft = True
    ws.freeze_panes = "C6"

    title_fill = PatternFill("solid", fgColor="E2F0D9")
    brand_fill = PatternFill("solid", fgColor="17365D")
    group_fill = PatternFill("solid", fgColor="366092")
    section_fill = PatternFill("solid", fgColor="4472C4")
    flavor_fill = PatternFill("solid", fgColor="5B9BD5")
    total_fill = PatternFill("solid", fgColor="D9EAF7")

    header_font = Font(color="FFFFFF", bold=True)
    bold_font = Font(bold=True)
    thin = Side(style="thin", color="D9E1F2")

    for cell in ws[1]:
        cell.fill = title_fill
        cell.font = Font(bold=True, size=12)
        cell.alignment = Alignment(
            horizontal="right",
            vertical="center",
        )

    for row_idx, fill in (
        (2, brand_fill),
        (3, group_fill),
        (4, section_fill),
        (5, flavor_fill),
    ):
        for cell in ws[row_idx]:
            cell.fill = fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
                readingOrder=2,
            )
            cell.border = Border(
                left=thin,
                right=thin,
                top=thin,
                bottom=thin,
            )

    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 25
    ws.row_dimensions[3].height = 25
    ws.row_dimensions[4].height = 25
    ws.row_dimensions[5].height = 48

    if max_row >= 6:
        for cell in ws[max_row]:
            cell.fill = total_fill
            cell.font = bold_font
            cell.border = Border(top=thin)

    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 16

    for col_idx in range(3, max_col + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 17

    for row in ws.iter_rows(
        min_row=6,
        max_row=max_row,
        min_col=1,
        max_col=max_col,
    ):
        for cell in row:
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
                readingOrder=2,
            )


def merge_same_values_in_row(ws, row, start_col, values, parent_keys=None):
    """
    مقادیر پشت سر هم را در یک ردیف Merge می‌کند.
    parent_keys باعث می‌شود Merge از مرز والد عبور نکند.
    """
    if not values:
        return

    block_start = start_col
    current_value = values[0]
    current_parent = parent_keys[0] if parent_keys else None

    for offset in range(1, len(values)):
        col = start_col + offset
        value = values[offset]
        parent = parent_keys[offset] if parent_keys else None

        if value != current_value or parent != current_parent:
            end_col = col - 1
            if end_col > block_start:
                ws.merge_cells(
                    start_row=row,
                    start_column=block_start,
                    end_row=row,
                    end_column=end_col,
                )

            block_start = col
            current_value = value
            current_parent = parent

    end_col = start_col + len(values) - 1

    if end_col > block_start:
        ws.merge_cells(
            start_row=row,
            start_column=block_start,
            end_row=row,
            end_column=end_col,
        )


def merge_hierarchical_headers(ws, output_keys):
    """
    هدر بدون Merge هم‌پوشان:

    ردیف 2: برند
    ردیف 3: گروه کالا
    ردیف 4: سکشن
    ردیف 5: طعم

    اگر یک گروه سکشن نداشته باشد، نام گروه یک‌بار در محدوده‌ی
    ردیف 3 تا 4 و تمام ستون‌های همان گروه Merge می‌شود.

    اگر گروه سکشن داشته باشد:
      - گروه فقط در ردیف 3 روی کل ستون‌های خودش Merge می‌شود.
      - هر سکشن در ردیف 4 روی تمام طعم‌های خودش Merge می‌شود.
    """
    # ستون‌های ثابت نماینده
    ws.merge_cells(
        start_row=2,
        start_column=1,
        end_row=5,
        end_column=1,
    )
    ws.cell(2, 1).value = "نام کامل نماینده"

    ws.merge_cells(
        start_row=2,
        start_column=2,
        end_row=5,
        end_column=2,
    )
    ws.cell(2, 2).value = "کد شرکت"

    if not output_keys:
        return

    # --------------------------------------------------------
    # 1) برندها: Merge افقی در ردیف 2
    # --------------------------------------------------------
    idx = 0

    while idx < len(output_keys):
        brand = output_keys[idx][0]
        block_end = idx

        while (
            block_end + 1 < len(output_keys)
            and output_keys[block_end + 1][0] == brand
        ):
            block_end += 1

        start_col = 3 + idx
        end_col = 3 + block_end

        if end_col > start_col:
            ws.merge_cells(
                start_row=2,
                start_column=start_col,
                end_row=2,
                end_column=end_col,
            )

        ws.cell(2, start_col).value = brand
        idx = block_end + 1

    # --------------------------------------------------------
    # 2) گروه کالا + سکشن
    #    این بخش از اول طراحی شده تا هیچ Merge هم‌پوشانی نداشته باشد.
    # --------------------------------------------------------
    idx = 0

    while idx < len(output_keys):
        brand = output_keys[idx][0]
        group = output_keys[idx][1]

        group_end = idx

        while (
            group_end + 1 < len(output_keys)
            and output_keys[group_end + 1][0] == brand
            and output_keys[group_end + 1][1] == group
        ):
            group_end += 1

        start_col = 3 + idx
        end_col = 3 + group_end

        group_sections = [
            normalize_text(output_keys[pos][2])
            for pos in range(idx, group_end + 1)
        ]

        has_any_section = any(group_sections)

        if not has_any_section:
            # هیچ سکشنی ندارد:
            # خود گروه روی دو ردیف 3 و 4 و تمام ستون‌هایش Merge می‌شود.
            ws.merge_cells(
                start_row=3,
                start_column=start_col,
                end_row=4,
                end_column=end_col,
            )
            ws.cell(3, start_col).value = group

        else:
            # گروه دارای سکشن است:
            # گروه فقط روی ردیف 3 Merge می‌شود.
            if end_col > start_col:
                ws.merge_cells(
                    start_row=3,
                    start_column=start_col,
                    end_row=3,
                    end_column=end_col,
                )

            ws.cell(3, start_col).value = group

            # هر سکشن روی تمام طعم‌های خودش Merge می‌شود.
            section_idx = idx

            while section_idx <= group_end:
                section = normalize_text(
                    output_keys[section_idx][2]
                )
                section_end = section_idx

                while (
                    section_end + 1 <= group_end
                    and normalize_text(
                        output_keys[section_end + 1][2]
                    ) == section
                ):
                    section_end += 1

                section_start_col = 3 + section_idx
                section_end_col = 3 + section_end

                if section:
                    if section_end_col > section_start_col:
                        ws.merge_cells(
                            start_row=4,
                            start_column=section_start_col,
                            end_row=4,
                            end_column=section_end_col,
                        )

                    ws.cell(4, section_start_col).value = section

                else:
                    # اگر در یک گروه دارای سکشن، بخشی سکشن خالی داشت،
                    # همان محدوده‌ی خالی را هم یکپارچه Merge می‌کنیم
                    # تا ظاهر هدر شکسته نشود.
                    if section_end_col > section_start_col:
                        ws.merge_cells(
                            start_row=4,
                            start_column=section_start_col,
                            end_row=4,
                            end_column=section_end_col,
                        )

                    ws.cell(4, section_start_col).value = ""

                section_idx = section_end + 1

        idx = group_end + 1


def build_output(
    output_path: Path,
    processed_results: list[dict],
    output_keys: list[tuple[str, str, str, str]],
    errors: list[dict],
):
    wb = Workbook()

    ws = wb.active
    ws.title = "گزارش موجودی کشوری لاین 2"

    jalali_date = today_jalali_string()

    title = (
        f"تاریخ گزارش: {jalali_date}   |   "
        "National Inventory Report"
    )

    total_columns = 2 + len(output_keys)

    ws.append(
        [title]
        + [""] * max(0, total_columns - 1)
    )

    if total_columns >= 2:
        ws.merge_cells(
            start_row=1,
            start_column=1,
            end_row=1,
            end_column=total_columns,
        )

    # هرم چهارسطحی
    ws.append(
        ["", ""]
        + [key[0] for key in output_keys]
    )
    ws.append(
        ["", ""]
        + [key[1] for key in output_keys]
    )
    ws.append(
        ["", ""]
        + [key[2] for key in output_keys]
    )
    ws.append(
        ["نام کامل نماینده", "کد شرکت"]
        + [key[3] for key in output_keys]
    )

    merge_hierarchical_headers(ws, output_keys)

    processed_results = sorted(
        processed_results,
        key=lambda item: persian_sort_key(
            item.get("company_name")
            or item["representative"]
        ),
    )

    for item in processed_results:
        totals = item["totals"]

        row = [
            item.get("company_name")
            or item["representative"],
            item.get("company_code", ""),
        ]

        for output_key in output_keys:
            row.append(
                int(totals.get(output_key, 0))
            )

        ws.append(row)

    total_row = ["مجموع(کارتن)", ""]

    for output_key in output_keys:
        total_row.append(
            sum(
                int(item["totals"].get(output_key, 0))
                for item in processed_results
            )
        )

    ws.append(total_row)

    last_row = ws.max_row
    last_col = ws.max_column

    style_output_sheet(
        ws,
        last_col,
        last_row,
    )

    if last_col >= 2 and last_row >= 6:
        ws.auto_filter.ref = (
            f"A5:"
            f"{get_column_letter(last_col)}"
            f"{max(5, last_row - 1)}"
        )

    # ------------------ شیت خطاها ------------------
    err_ws = wb.create_sheet("خطاها")
    err_ws.sheet_view.rightToLeft = True
    err_ws.freeze_panes = "A2"

    error_headers = [
        "نام فایل",
        "نام نماینده",
        "نام شیت",
        "ردیف فایل",
        "کد کالا",
        "نام کالا در فایل",
        "نوع خطا",
        "توضیح",
        "ردیف دیکشنری",
        "ستون مورد نیاز",
    ]
    err_ws.append(error_headers)

    if errors:
        for error in errors:
            err_ws.append(
                [
                    error.get(h, "")
                    for h in error_headers
                ]
            )
    else:
        err_ws.append(
            [
                "",
                "",
                "",
                "",
                "",
                "",
                "بدون خطا",
                "خطایی ثبت نشد.",
                "",
                "",
            ]
        )

    header_fill = PatternFill(
        "solid",
        fgColor="9C0006",
    )
    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    for cell in err_ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
            readingOrder=2,
        )

    widths = {
        "A": 28,
        "B": 28,
        "C": 20,
        "D": 14,
        "E": 18,
        "F": 42,
        "G": 28,
        "H": 65,
        "I": 16,
        "J": 25,
    }

    for letter, width in widths.items():
        err_ws.column_dimensions[letter].width = width

    for row in err_ws.iter_rows(
        min_row=2,
        max_row=err_ws.max_row,
    ):
        for cell in row:
            cell.alignment = Alignment(
                horizontal="right",
                vertical="top",
                wrap_text=True,
                readingOrder=2,
            )

    try:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
    except Exception:
        pass

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )
    wb.save(output_path)


# ============================================================
# اجرای اصلی
# ============================================================

def main():
    print("=" * 70)
    print("شروع تبدیل موجودی نمایندگان به کارتن مادر")
    print("=" * 70)
    print("خروجی: برند -> گروه کالا -> سکشن -> طعم | نام شیت: گزارش موجودی کشوری لاین 2")
    print("قاعده تجمیع: ردیف‌های تکراری جمع می‌شوند و در تعریف‌های جایگزین، سطح 3 سپس 2 و سپس 1 اولویت دارد.")

    if not INPUT_FOLDER.exists():
        INPUT_FOLDER.mkdir(parents=True, exist_ok=True)
        print(f"\nپوشه ورودی ساخته شد:\n{INPUT_FOLDER}")
        print("فایل‌های نماینده را داخل این پوشه قرار بده و برنامه را دوباره اجرا کن.")
        return

    print("\nدر حال خواندن دیکشنری...", flush=True)

    try:
        dictionary_file = resolve_dictionary_file()
        print(
            f"دیکشنری انتخاب‌شده: {dictionary_file.name}",
            flush=True,
        )

        dictionary_data = load_dictionary(dictionary_file)
        print(
            f"دیکشنری خوانده شد | نماینده‌ها: {len(dictionary_data['representatives'])} "
            f"| ستون‌های خروجی: {len(dictionary_data['output_keys'])}",
            flush=True,
        )
    except Exception as exc:
        print("\nخطا در خواندن دیکشنری:")
        print(exc)
        sys.exit(1)

    input_files = sorted(
        [
            path
            for path in INPUT_FOLDER.iterdir()
            if path.is_file()
            and path.suffix.lower() in {".xlsx", ".xlsm"}
            and not path.name.startswith("~$")
            and path.resolve() != OUTPUT_FILE.resolve()
        ],
        key=lambda p: p.name,
    )

    if not input_files:
        print(f"\nهیچ فایل xlsx/xlsm در پوشه زیر پیدا نشد:\n{INPUT_FOLDER}")
        return

    errors = []
    processed_results = []

    print(f"\nتعداد فایل‌های ورودی: {len(input_files)}")

    for index, path in enumerate(input_files, start=1):
        print(f"[{index}/{len(input_files)}] {path.name}", flush=True)

        result = process_inventory_file(
            path=path,
            dictionary_data=dictionary_data,
            errors=errors,
        )

        if result is not None:
            processed_results.append(result)

    build_output(
        output_path=OUTPUT_FILE,
        processed_results=processed_results,
        output_keys=dictionary_data["output_keys"],
        errors=errors,
    )

    print("\n" + "=" * 70)
    print("پایان پردازش")
    print(f"نماینده‌های پردازش‌شده: {len(processed_results)}")
    print(f"تعداد خطاهای ثبت‌شده: {len(errors)}")
    print(f"فایل خروجی:\n{OUTPUT_FILE}")
    print("=" * 70)


if __name__ == "__main__":
    main()
